#!/usr/bin/env python3
"""Vuelca el texto de un pptx (por lámina), docx (párrafos y tablas), md o txt.

Uso: extraer_texto.py archivo [salida.txt]
Sirve para pasar el contenido a los agentes de revisión y a los otros scripts.
"""
import sys
from pathlib import Path


def texto_pptx(ruta):
    from pptx import Presentation
    out = []
    for i, sl in enumerate(Presentation(ruta).slides, 1):
        out.append(f"=== Lámina {i}")
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    out.append(" | ".join(c.text for c in r.cells))
    return "\n".join(out)


def texto_docx(ruta):
    import docx
    d = docx.Document(ruta)
    out = [p.text for p in d.paragraphs if p.text.strip()]
    out.append("--- TABLAS ---")
    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join(c.text for c in r.cells))
    return "\n".join(out)


def extraer(ruta):
    ext = Path(ruta).suffix.lower()
    if ext == ".pptx":
        return texto_pptx(ruta)
    if ext == ".docx":
        return texto_docx(ruta)
    return Path(ruta).read_text(encoding="utf-8")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    t = extraer(sys.argv[1])
    if len(sys.argv) > 2:
        Path(sys.argv[2]).write_text(t, encoding="utf-8")
        print(f"Texto guardado en {sys.argv[2]} ({len(t.split())} palabras)")
    else:
        print(t)
